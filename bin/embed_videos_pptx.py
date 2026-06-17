#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import re
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu


ROOT = Path(__file__).resolve().parent.parent
SLIDES_MD = ROOT / "slides" / "tropical-ruby-2026.md"


@dataclass(frozen=True)
class Placement:
    left: float
    top: float
    width: float
    height: float


@dataclass(frozen=True)
class VideoSpec:
    movie: Path
    poster: Path
    mime_type: str
    placement: Placement


VIDEO_SPECS: dict[str, VideoSpec] = {
    "asamiarts-tracing": VideoSpec(
        movie=ROOT / "assets" / "asamiarts tracing.mp4",
        poster=ROOT / "assets" / "asamiarts tracing.jpg",
        mime_type="video/mp4",
        placement=Placement(left=0.468, top=0.258, width=0.448, height=0.448 * (9 / 16) * (13.333 / 7.5)),
    ),
    "tracing-hidden-layer": VideoSpec(
        movie=ROOT / "assets" / "tracing, hidden layer vertical.mp4",
        poster=ROOT / "assets" / "tracing, hidden layer vertical.jpg",
        mime_type="video/mp4",
        placement=Placement(left=0.714, top=0.148, width=0.223, height=0.804),
    ),
}


def emu(value: float, full: int) -> Emu:
    return Emu(int(full * value))


def find_marked_slides() -> dict[str, int]:
    text = SLIDES_MD.read_text()
    slides = text.split("\n---\n")
    # Marp front matter lives before the first thematic break and is not a real slide.
    if slides and slides[0].lstrip().startswith("---\nmarp:"):
        slides = slides[1:]
    markers: dict[str, int] = {}
    for index, slide in enumerate(slides, 1):
        matches = re.findall(r"<!--\s*pptx-video:\s*([a-z0-9-]+)\s*-->", slide)
        for key in matches:
            if key in markers:
                raise SystemExit(f"Duplicate pptx-video marker: {key}")
            markers[key] = index
    return markers


def main(argv: list[str]) -> int:
    if len(argv) not in (2, 3):
        print("usage: embed_videos_pptx.py INPUT_PPTX [OUTPUT_PPTX]", file=sys.stderr)
        return 2

    input_pptx = Path(argv[1]).resolve()
    output_pptx = Path(argv[2]).resolve() if len(argv) == 3 else input_pptx

    if not input_pptx.exists():
        raise SystemExit(f"Input PPTX not found: {input_pptx}")

    markers = find_marked_slides()
    missing_markers = sorted(set(VIDEO_SPECS) - set(markers))
    if missing_markers:
        raise SystemExit(f"Missing pptx-video markers in slides markdown: {', '.join(missing_markers)}")

    for key, spec in VIDEO_SPECS.items():
        if not spec.movie.exists():
            raise SystemExit(f"Missing movie asset for {key}: {spec.movie}")
        if not spec.poster.exists():
            raise SystemExit(f"Missing poster asset for {key}: {spec.poster}")

    if output_pptx != input_pptx:
        output_pptx.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(input_pptx, output_pptx)

    prs = Presentation(str(output_pptx))
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    for key, spec in VIDEO_SPECS.items():
        slide_no = markers[key]
        slide = prs.slides[slide_no - 1]

        # Find the picture Marp already placed for the poster image.
        # When present, reuse its position/size and remove it so the video
        # sits exactly where the placeholder was — no overlap, no resize.
        poster_pic = find_poster_picture(slide, spec.poster)
        if poster_pic is not None:
            left = poster_pic.left
            top = poster_pic.top
            width = poster_pic.width
            height = poster_pic.height
            remove_shape(poster_pic)
        else:
            place = spec.placement
            left = emu(place.left, slide_width)
            top = emu(place.top, slide_height)
            width = emu(place.width, slide_width)
            height = emu(place.height, slide_height)

        movie = slide.shapes.add_movie(
            str(spec.movie),
            left,
            top,
            width,
            height,
            poster_frame_image=str(spec.poster),
            mime_type=spec.mime_type,
        )
        enable_autoplay_and_loop(slide, movie.shape_id)

    prs.save(str(output_pptx))
    print(f"Embedded {len(VIDEO_SPECS)} videos into {output_pptx}")
    return 0


def find_poster_picture(slide, poster_path: Path):
    """Locate the picture shape on `slide` whose embedded image matches `poster_path`.

    Marp renders an HTML `<video poster="...">` (or the `<img>`-only fallback)
    as a regular picture shape in the PPTX. We match by SHA1 of the raw image
    blob — exact same bytes as the poster file on disk.
    """
    poster_sha = hashlib.sha1(poster_path.read_bytes()).hexdigest()
    for shape in slide.shapes:
        if not getattr(shape, "image", None):
            continue
        try:
            if shape.image.sha1 == poster_sha:
                return shape
        except (AttributeError, ValueError):
            continue
    return None


def remove_shape(shape) -> None:
    """Remove `shape` from its parent (PPTX has no first-class delete API)."""
    sp = shape._element
    sp.getparent().remove(sp)


def enable_autoplay_and_loop(slide, shape_id: int) -> None:
    """Flip the embedded movie timing from click-to-play to autoplay+loop."""
    root = slide.part._element
    videos = root.xpath(f".//p:video[p:cMediaNode/p:tgtEl/p:spTgt[@spid='{shape_id}']]")
    if not videos:
        raise SystemExit(f"Unable to find timing node for video shape {shape_id}")

    video = videos[0]
    ctn = video.xpath("./p:cMediaNode/p:cTn")[0]
    ctn.set("repeatCount", "indefinite")

    cond = video.xpath("./p:cMediaNode/p:cTn/p:stCondLst/p:cond")[0]
    cond.set("delay", "0")


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
